from datetime import timedelta

from django.urls import reverse
from django.utils import timezone

from ncs.models import NC
from ncs.test_workflow import WorkflowTestBase

from .stats import build_dashboard

S = NC.Status
TODAY = timezone.localdate


class DashboardTestBase(WorkflowTestBase):
    """Fixture NC (IT -> Finance, pending) plus a few more in other states."""

    def setUp(self):
        past = TODAY() - timedelta(days=5)
        self.overdue = self.new_nc(S.IN_PROGRESS, target_date=past, description="Late backup")
        self.closed = self.new_nc(S.CLOSED, receiving_department=self.hr,
                                  closed_at=timezone.now(), description="Closed one")
        self.not_valid = self.new_nc(S.NOT_VALID, receiving_department=self.hr)

    def ids(self, response):
        return {nc.nc_id for nc in response.context["page"]}


class RegisterTests(DashboardTestBase):
    url = reverse("dashboard:register")

    def get(self, user, **params):
        self.client.force_login(user)
        return self.client.get(self.url, params)

    def test_cards_by_default_table_on_request(self):
        cards = self.get(self.manager)
        self.assertEqual(cards.context["view_mode"], "cards")
        self.assertContains(cards, "nc-card")
        table = self.get(self.manager, view="table", status="CLOSED")
        self.assertEqual(table.context["view_mode"], "table")
        self.assertEqual(self.ids(table), {self.closed.nc_id})  # filters still apply
        self.assertFalse(self.get(self.manager, view="table").context["filters_used"])

    def test_quick_filter_counts(self):
        chips = {c["label"]: c for c in self.get(self.manager).context["quick_filters"]}
        counts = {label: chips[label]["count"] for label in ("All", "Open", "Overdue", "Closed")}
        self.assertEqual(counts, {"All": 4, "Open": 2, "Overdue": 1, "Closed": 1})
        self.assertTrue(chips["All"]["active"])
        # Counts follow the other filters, e.g. receiving department = HR
        chips = {c["label"]: c for c in self.get(self.manager, receiving_department=self.hr.pk)
                 .context["quick_filters"]}
        self.assertEqual((chips["All"]["count"], chips["Open"]["count"]), (2, 0))

    def test_most_urgent_first(self):
        page = self.get(self.manager, sort="urgent").context["page"]
        self.assertEqual(page[0], self.overdue)  # its target date has already passed

    def test_your_action_tag(self):
        # Finance HoD must validate self.nc; the Action Owner is working on self.overdue
        self.assertEqual(self.get(self.finance_hod).context["my_action_ids"], {self.nc.pk})
        self.assertEqual(self.get(self.owner).context["my_action_ids"], {self.overdue.pk})
        self.assertEqual(self.get(self.management).context["my_action_ids"], set())

    def test_hod_sees_whole_register(self):  # FR-30
        self.assertEqual(self.get(self.it_hod).context["total"], NC.objects.count())

    def test_department_staff_see_only_their_department(self):
        response = self.get(self.hr_staff)
        self.assertEqual(self.ids(response), {self.closed.nc_id, self.not_valid.nc_id})

    def test_filters(self):  # FR-32
        cases = [
            ({"status": "CLOSED"}, {self.closed.nc_id}),
            ({"overdue": "yes"}, {self.overdue.nc_id}),
            ({"receiving_department": self.hr.pk}, {self.closed.nc_id, self.not_valid.nc_id}),
            ({"q": "backup"}, {self.overdue.nc_id}),
            ({"q": self.closed.nc_id}, {self.closed.nc_id}),
            ({"action_owner": self.owner.pk, "status": "open"}, {self.nc.nc_id, self.overdue.nc_id}),
        ]
        for params, expected in cases:
            with self.subTest(params=params):
                self.assertEqual(self.ids(self.get(self.manager, **params)), expected)

    def test_open_excludes_closed_and_not_valid(self):  # FR-15
        ids = self.ids(self.get(self.manager, status="open"))
        self.assertNotIn(self.closed.nc_id, ids)
        self.assertNotIn(self.not_valid.nc_id, ids)

    def test_date_range_and_bad_input(self):
        future = (TODAY() + timedelta(days=1)).isoformat()
        self.assertEqual(self.get(self.manager, date_from=future).context["total"], 0)
        response = self.get(self.manager, date_from="not-a-date")
        self.assertEqual(response.status_code, 200)  # ignored, not a crash

    def test_rows_coloured(self):  # UI-04
        self.assertContains(self.get(self.manager), "rag-red")


class ExportTests(DashboardTestBase):
    def test_excel_has_filtered_rows(self):  # FR-36
        from io import BytesIO

        from openpyxl import load_workbook

        self.client.force_login(self.manager)
        response = self.client.get(reverse("dashboard:export_excel"), {"status": "open"})
        self.assertEqual(response["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet["A4"].value, "NC ID")
        ids = {row[0] for row in sheet.iter_rows(min_row=5, values_only=True)}
        self.assertEqual(ids, {self.nc.nc_id, self.overdue.nc_id})
        self.assertIn("Status: All open", sheet["A1"].value)

    def test_excel_respects_department_view(self):
        from io import BytesIO

        from openpyxl import load_workbook

        self.client.force_login(self.hr_staff)
        sheet = load_workbook(BytesIO(self.client.get(reverse("dashboard:export_excel")).content)).active
        ids = {row[0] for row in sheet.iter_rows(min_row=5, values_only=True)}
        self.assertEqual(ids, {self.closed.nc_id, self.not_valid.nc_id})

    def test_pdf_or_printable_fallback(self):
        from unittest import mock

        from . import exports

        self.client.force_login(self.manager)
        with mock.patch.object(exports, "pdf_available", return_value=False):
            response = self.client.get(reverse("dashboard:export_pdf"), {"status": "CLOSED"})
        self.assertContains(response, "print-ready page")
        self.assertContains(response, self.closed.nc_id)
        self.assertNotContains(response, self.overdue.nc_id)

        if exports.pdf_available():  # only on machines with WeasyPrint's libraries
            response = self.client.get(reverse("dashboard:export_pdf"))
            self.assertEqual(response["Content-Type"], "application/pdf")
            self.assertTrue(response.content.startswith(b"%PDF"))


class DashboardTests(DashboardTestBase):
    def test_counts(self):  # FR-33
        counts = build_dashboard(NC.objects.all())["counts"]
        self.assertEqual(counts["total"], 4)
        self.assertEqual(counts["open"], 2)  # pending + in progress
        self.assertEqual(counts["overdue"], 1)
        self.assertEqual(counts["closed"], 1)
        self.assertEqual(counts["not_valid"], 1)

    def test_breakdowns_exclude_not_valid(self):  # FR-34, FR-15
        data = build_dashboard(NC.objects.all())
        self.assertEqual(dict(data["by_receiving"]), {"Finance": 2, "HR": 1})
        self.assertEqual(dict(data["by_raising"]), {"IT": 3})

    def test_trend_and_closure(self):  # FR-35
        data = build_dashboard(NC.objects.all())
        self.assertEqual(len(data["trend"]["labels"]), 12)
        self.assertEqual(data["trend"]["raised"][-1], 4)   # all logged this month
        self.assertEqual(data["trend"]["closed"][-1], 1)
        self.assertEqual(sum(n for _, n in data["closure"]["ageing"]), 1)
        self.assertEqual(data["closure"]["average_days"], 0)

    def test_page_for_every_role(self):
        for user in [self.manager, self.management, self.it_hod, self.owner, self.hr_staff]:
            with self.subTest(user=user.username):
                self.client.force_login(user)
                response = self.client.get(reverse("dashboard:dashboard"))
                self.assertContains(response, 'id="chart-data"')

    def test_department_view_for_staff(self):
        self.client.force_login(self.hr_staff)
        response = self.client.get(reverse("dashboard:dashboard"))
        tiles = {label: value for label, value, _, _ in response.context["tiles"]}
        self.assertEqual(tiles["Total logged"], 2)  # only HR's NCs
