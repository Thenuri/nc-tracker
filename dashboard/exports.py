"""Register exports (FR-36). Both use the same filtered list as the screen."""

from functools import cache
from io import BytesIO

from django.conf import settings
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# (column heading, function that returns the cell value, column width)
COLUMNS = [
    ("NC ID", lambda nc: nc.nc_id, 13),
    ("Status", lambda nc: nc.get_status_display(), 20),
    ("Overdue", lambda nc: f"Yes ({nc.days_overdue} days)" if nc.is_overdue else "No", 14),
    ("Raising department", lambda nc: str(nc.raising_department), 20),
    ("Receiving department", lambda nc: str(nc.receiving_department), 20),
    ("Area / process", lambda nc: f"{nc.process}{' – ' + nc.process_other if nc.process_other else ''}", 24),
    ("Description", lambda nc: nc.description, 50),
    ("Date identified", lambda nc: nc.date_identified, 14),
    ("Source", lambda nc: str(nc.source), 18),
    ("Identified by", lambda nc: nc.identified_by, 20),
    ("Severity", lambda nc: nc.get_severity_display(), 10),
    ("Logged", lambda nc: timezone.localtime(nc.logged_at).replace(tzinfo=None), 17),
    ("Validation deadline", lambda nc: nc.validation_deadline, 14),
    ("Dispute reason", lambda nc: nc.dispute_reason, 30),
    ("Dispute decision", lambda nc: nc.dispute_rationale, 30),
    ("Root cause", lambda nc: nc.root_cause, 35),
    ("Corrective action", lambda nc: nc.corrective_action, 35),
    ("Action Owner", lambda nc: str(nc.action_owner or ""), 20),
    ("Target date", lambda nc: nc.target_date, 14),
    ("Original target date", lambda nc: nc.original_target_date, 14),
    ("Completion date", lambda nc: nc.completion_date, 14),
    ("Evidence", lambda nc: evidence_text(nc), 40),
    ("Verification", lambda nc: nc.get_verification_outcome_display(), 14),
    ("Verification comments", lambda nc: nc.verification_comments, 30),
    ("Closed", lambda nc: timezone.localtime(nc.closed_at).replace(tzinfo=None) if nc.closed_at else None, 17),
    ("Days open", lambda nc: nc.days_open, 10),
]


def evidence_text(nc):
    """Evidence descriptions with links, one per line (auditors' export, Section 9)."""
    lines = []
    for item in nc.evidence.all():
        download = settings.SITE_URL + reverse("ncs:evidence_download", args=[item.pk])
        where = item.link or (download if item.file else "")
        lines.append(f"{item.description}: {where}")
    return "\n".join(lines)


def register_excel(ncs, title):
    """An .xlsx file (bytes) with one row per NC."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Register"

    sheet.append([title])
    sheet["A1"].font = Font(bold=True, size=13)
    sheet.append([f"Exported {timezone.localtime():%d %b %Y %H:%M}"
                  + (" · PROTOTYPE – sample data" if settings.PROTOTYPE_MODE else "")])
    sheet.append([])

    header_row = 4
    sheet.append([heading for heading, _, _ in COLUMNS])
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0D6EFD")

    for nc in ncs.prefetch_related("evidence"):
        sheet.append([value(nc) for _, value, _ in COLUMNS])

    for index, (_, _, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if hasattr(cell.value, "year"):
                cell.number_format = "dd mmm yyyy" if not hasattr(cell.value, "hour") else "dd mmm yyyy hh:mm"

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=2)  # keep header + NC ID visible
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(COLUMNS))}{sheet.max_row}"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@cache  # check once per run; a failed import prints a long warning each time
def pdf_available():
    """WeasyPrint needs system libraries (GTK/Pango). Check they are present."""
    try:
        import weasyprint  # noqa: F401
    except (ImportError, OSError):
        return False
    return True


def register_html(request, ncs, title, filters_text):
    """Landscape, print-ready register. Used for the PDF and as its fallback."""
    return render_to_string("dashboard/register_print.html", {
        "ncs": ncs, "title": title, "filters_text": filters_text,
        "exported_at": timezone.localtime(), "pdf_fallback": not pdf_available(),
    }, request=request)


def register_pdf(request, ncs, title, filters_text):
    """PDF bytes, or None if WeasyPrint can't run on this machine."""
    if not pdf_available():
        return None
    import weasyprint

    html = register_html(request, ncs, title, filters_text)
    return weasyprint.HTML(string=html, base_url=request.build_absolute_uri("/")).write_pdf()
